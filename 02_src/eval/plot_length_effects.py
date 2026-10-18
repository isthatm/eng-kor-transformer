"""Plot length diagnostics from a prediction XLSX into its sibling artifacts/ folder.

Install dependencies: python -m pip install numpy matplotlib openpyxl
Run: python plot_length_effects.py INPUT.xlsx
Scores are means of sentence-level chrF++, not recomputed corpus chrF++.
"""

import argparse
from pathlib import Path
from zipfile import BadZipFile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import StrMethodFormatter
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


BIN_WIDTH = 5
BOOTSTRAP_SAMPLES = 5000
SEED = 42
BLUE = "#2878C8"
ORANGE = "#CC701F"


def load_predictions(path: Path) -> tuple[np.ndarray, str]:
    """Read the first sheet with prediction columns; reject malformed observations."""
    required = ("src_tokens", "chrf++", "hyp_tokens", "ref_tokens")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            if not all(name in header for name in required):
                continue
            if any(header.count(name) != 1 for name in required):
                raise ValueError(f"Sheet {sheet.title!r} has duplicate prediction columns.")
            indices = [header.index(name) for name in required]
            predictions = []
            for row_number, row in enumerate(rows, 2):
                if not any(value is not None for value in row):
                    continue
                try:
                    raw = [row[index] for index in indices]
                    if any(isinstance(value, bool) for value in raw):
                        raise ValueError
                    source, score, hypothesis, reference = map(float, raw)
                    if not np.isfinite([source, score, hypothesis, reference]).all():
                        raise ValueError
                    if not (0 <= score <= 100 and source >= 1):
                        raise ValueError
                    if any(value < 0 or not value.is_integer()
                           for value in (source, hypothesis, reference)):
                        raise ValueError
                except (IndexError, TypeError, ValueError) as error:
                    raise ValueError(
                        f"Invalid prediction at {sheet.title!r}, Excel row {row_number}: "
                        "token counts must be nonnegative integers (source > 0), "
                        "and chrF++ must be finite and between 0 and 100."
                    ) from error
                ratio = hypothesis / reference if reference else np.nan
                predictions.append((source, score, ratio))
            if not predictions:
                raise ValueError(f"Sheet {sheet.title!r} has no prediction rows.")
            return np.asarray(predictions), sheet.title
    finally:
        workbook.close()
    raise ValueError(f"No sheet contains all required columns: {', '.join(required)}.")


def bootstrap_mean_ci(scores: np.ndarray, rng: np.random.Generator) -> np.ndarray:
    # shortcut: resample sentences independently; use document clusters when IDs are available.
    means = np.empty(BOOTSTRAP_SAMPLES)
    for start in range(0, BOOTSTRAP_SAMPLES, 256):
        stop = min(start + 256, BOOTSTRAP_SAMPLES)
        indices = rng.integers(0, len(scores), size=(stop - start, len(scores)))
        means[start:stop] = scores[indices].mean(axis=1)
    return np.percentile(means, [2.5, 97.5])


def length_statistics(data: np.ndarray) -> dict:
    """Use inclusive integer bins 1-5, 6-10, ...; preserve empty bins as gaps."""
    membership = ((data[:, 0] - 1) // BIN_WIDTH).astype(int)
    bins = np.arange(membership.min(), membership.max() + 1)
    counts = np.zeros(len(bins), dtype=int)
    means = np.full(len(bins), np.nan)
    quantiles = np.full((len(bins), 5), np.nan)
    intervals = np.full((len(bins), 2), np.nan)
    rng = np.random.default_rng(SEED)
    for i, bin_id in enumerate(bins):
        scores = data[membership == bin_id, 1]
        counts[i] = len(scores)
        if len(scores):
            means[i] = scores.mean()
            quantiles[i] = np.percentile(scores, [10, 25, 50, 75, 90])
            intervals[i] = bootstrap_mean_ci(scores, rng)
    return {
        "centers": bins * BIN_WIDTH + (BIN_WIDTH + 1) / 2,
        "labels": [f"{b * BIN_WIDTH + 1}\u2013{(b + 1) * BIN_WIDTH}" for b in bins],
        "counts": counts, "means": means, "quantiles": quantiles, "ci": intervals,
    }


def create_plots(input_xlsx: Path) -> list[Path]:
    data, sheet = load_predictions(input_xlsx)
    valid_ratio = np.isfinite(data[:, 2])
    ratio_data = data[valid_ratio]
    if not len(ratio_data):
        raise ValueError("No positive reference-token counts; length ratios are undefined.")
    stats = length_statistics(data)
    output_dir = input_xlsx.parent / "artifacts"
    output_dir.mkdir(exist_ok=True)
    paths = [output_dir / name for name in (
        "01_chrf_by_source_length.png",
        "02_chrf_by_length_ratio.png",
        "03_chrf_distribution_by_source_length.png",
    )]
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 11,
        "axes.spines.top": False, "axes.spines.right": False,
        "axes.axisbelow": True, "grid.color": "#DFE4EA", "grid.linewidth": 0.8,
        "axes.titlepad": 14, "savefig.facecolor": "white",
    })
    x, means, counts = stats["centers"], stats["means"], stats["counts"]
    lower, upper = stats["ci"].T

    fig, ax = plt.subplots(figsize=(13.5, 8), layout="constrained")
    count_ax = ax.twinx()
    ax.set_zorder(count_ax.get_zorder() + 1)
    ax.patch.set_visible(False)
    count_ax.spines["right"].set_visible(True)
    fig.suptitle("Translation quality by source length", fontsize=19, fontweight="bold")
    ax.set_title(
        f"All {len(data):,} sentences. Five-token bins. "
        f"95% percentile CIs: {BOOTSTRAP_SAMPLES:,} sentence bootstrap resamples (seed {SEED}).",
        fontsize=10, color="#59636E",
    )
    ax.plot(x, means, "o-", color=BLUE, linewidth=2.2, label="Mean chrF++")
    ax.errorbar(
        x, (lower + upper) / 2, yerr=(upper - lower) / 2,
        fmt="none", ecolor=BLUE, capsize=5, linewidth=1.5, label="95% bootstrap CI",
    )
    ax.set_ylabel("Mean sentence chrF++", color=BLUE)
    ax.tick_params(axis="y", colors=BLUE)
    score_min = np.nanmin(np.r_[lower, means])
    score_max = np.nanmax(np.r_[upper, means])
    ax.set_ylim(max(0, 5 * np.floor(score_min / 5) - 5),
                min(100, 5 * np.ceil(score_max / 5) + 5))
    ax.grid(axis="y")
    count_ax.bar(x, counts, width=3.5, color="#A8BFD7", alpha=0.45,
                 label="Sentences per bin")
    for center, count in zip(x, counts):
        count_ax.annotate(f"{count:,}", (center, count), xytext=(0, 4),
                          textcoords="offset points", ha="center", fontsize=9,
                          color="#59636E")
    count_ax.set_ylim(0, max(counts) * 1.6)
    count_ax.set_ylabel("Sentences per bin", color="#59636E")
    count_ax.tick_params(axis="y", colors="#59636E")
    count_ax.yaxis.set_major_formatter(StrMethodFormatter("{x:,.0f}"))
    ax.set_xticks(x, stats["labels"])
    ax.set_xlabel("Source length (tokens)")
    ax.set_xlim(x[0] - BIN_WIDTH / 2, x[-1] + BIN_WIDTH / 2)
    handles, labels = ax.get_legend_handles_labels()
    count_handles, count_labels = count_ax.get_legend_handles_labels()
    ax.legend(handles + count_handles, labels + count_labels, frameon=False)
    fig.savefig(paths[0], dpi=180)
    plt.close(fig)

    edges = np.array([0, 0.5, 0.7, 0.9, 1.1, 1.3, 1.5, np.inf])
    ratio_labels = ["<0.5", "0.5\u2013<0.7", "0.7\u2013<0.9", "0.9\u2013<1.1",
                    "1.1\u2013<1.3", "1.3\u2013<1.5", "\u22651.5"]
    groups = np.searchsorted(edges, ratio_data[:, 2], side="right") - 1
    ratio_counts = np.bincount(groups, minlength=len(ratio_labels))
    ratio_means = np.array([
        ratio_data[groups == i, 1].mean() if count else np.nan
        for i, count in enumerate(ratio_counts)
    ])
    fig, ax = plt.subplots(figsize=(13.5, 7), layout="constrained")
    fig.suptitle("Output length ratio and translation quality", fontsize=19, fontweight="bold")
    ax.set_title(
        "Ratio = hypothesis tokens / reference tokens. "
        f"n = {len(ratio_data):,}; undefined ratios excluded = {(~valid_ratio).sum():,}.\n"
        "Bands are 0.2 wide around 1.0, with pooled tails below 0.5 and at least 1.5.",
        fontsize=10, color="#59636E",
    )
    positions = np.arange(len(ratio_labels))
    ax.bar(positions, ratio_means, width=0.65,
           color=[BLUE if i == 3 else "#A8BFD7" for i in positions])
    for i, mean in enumerate(ratio_means):
        if np.isfinite(mean):
            ax.annotate(f"{mean:.1f}", (i, mean), xytext=(0, 6),
                        textcoords="offset points", ha="center", fontweight="bold")
    ax.set_xticks(positions, [f"{label}\nn = {count:,}"
                             for label, count in zip(ratio_labels, ratio_counts)])
    ax.set_xlabel("Hypothesis / reference token-length ratio")
    ax.set_ylabel("Mean sentence chrF++")
    ax.set_ylim(0, 100)
    ax.grid(axis="y")
    fig.savefig(paths[1], dpi=180)
    plt.close(fig)

    p10, p25, median, p75, p90 = stats["quantiles"].T
    fig, ax = plt.subplots(figsize=(13.5, 8), layout="constrained")
    fig.suptitle("Translation score distribution by source length",
                 fontsize=19, fontweight="bold")
    ax.set_title(f"All {len(data):,} sentences. Five-token bins. "
                 "Shaded bands describe score distributions within each bin.",
                 fontsize=10, color="#59636E")
    ax.fill_between(x, p10, p90, color=BLUE, alpha=0.13,
                    linewidth=0, label="10th\u201390th percentiles (middle 80%)")
    ax.fill_between(x, p25, p75, color=BLUE, alpha=0.3,
                    linewidth=0, label="25th\u201375th percentiles (middle 50%)")
    ax.plot(x, median, "o-", color=BLUE, linewidth=2.5, markersize=7,
            markeredgecolor="white", markeredgewidth=1.5, label="Median")
    ax.plot(x, means, "D--", color=ORANGE, linewidth=1.8, markersize=5, label="Mean")
    ax.set_xticks(x, stats["labels"])
    ax.set_xlim(x[0] - BIN_WIDTH / 2, x[-1] + BIN_WIDTH / 2)
    ax.set_ylim(0, 100)
    ax.set_yticks([0, 25, 50, 75, 100])
    ax.set_xlabel("Source length (tokens)")
    ax.set_ylabel("Sentence-level chrF++")
    ax.grid(axis="y")
    ax.legend(frameon=False, loc="upper right", fontsize=10)
    fig.savefig(paths[2], dpi=180)
    plt.close(fig)

    print(f"Read {len(data):,} predictions from sheet {sheet!r}.")
    print(f"Source-token range: {int(data[:, 0].min())}\u2013{int(data[:, 0].max())}.")
    print("Source bin | Count | Mean chrF++ | 95% bootstrap CI")
    for label, count, mean, (lo, hi) in zip(stats["labels"], counts, means, stats["ci"]):
        print(f"{label:>10} | {count:5,d} | {mean:11.2f} | [{lo:.2f}, {hi:.2f}]")
    for path in paths:
        print(f"Saved {path}")
    return paths


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_xlsx", type=Path, help="Prediction workbook to analyze.")
    args = parser.parse_args()
    try:
        create_plots(args.input_xlsx.resolve())
    except (OSError, ValueError, BadZipFile, InvalidFileException) as error:
        parser.error(str(error))


if __name__ == "__main__":
    main()
